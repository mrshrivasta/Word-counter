from flask import Blueprint, request, jsonify, send_file
from app.analysis.metrics import get_basic_stats
from app.analysis.readability import get_readability_stats
from app.analysis.seo import get_seo_analysis
from werkzeug.utils import secure_filename
import os
from docx import Document
from reportlab.pdfgen import canvas
from io import BytesIO

api_bp = Blueprint('api', __name__)

@api_bp.route('/analyze', methods=['POST'])
def analyze():
    data = request.json
    text = data.get('text', '')

    seo_data = get_seo_analysis(text)
    return jsonify({
        'stats': get_basic_stats(text),
        'readability': get_readability_stats(text),
        'seo': seo_data,
        'advanced': {
            'geo_score': seo_data.get('geo_score', 0),
            'aeo_score': seo_data.get('aeo_score', 0),
            'aio_score': seo_data.get('aio_score', 0),
            'sxo_score': seo_data.get('sxo_score', 0)
        }
    })

@api_bp.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400

    filename = secure_filename(file.filename)
    upload_folder = 'static/uploads'
    if not os.path.exists(upload_folder):
        os.makedirs(upload_folder)
    filepath = os.path.join(upload_folder, filename)
    file.save(filepath)

    content = ""
    if filename.endswith('.txt'):
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
    elif filename.endswith('.docx'):
        doc = Document(filepath)
        content = "\n".join([para.text for para in doc.paragraphs])

    return jsonify({'content': content})

@api_bp.route('/export/<format>', methods=['POST'])
def export_file(format):
    # ... logic from app.py ...
    data = request.json
    text = data.get('text', '')

    if format == 'txt':
        buffer = BytesIO()
        buffer.write(text.encode('utf-8'))
        buffer.seek(0)
        return send_file(buffer, as_attachment=True, download_name='export.txt', mimetype='text/plain')

    elif format == 'pdf':
        buffer = BytesIO()
        p = canvas.Canvas(buffer)
        p.drawString(100, 800, "Text Export")
        y = 780
        for line in text.split('\n'):
            p.drawString(100, y, line)
            y -= 15
            if y < 50:
                p.showPage()
                y = 800
        p.save()
        buffer.seek(0)
        return send_file(buffer, as_attachment=True, download_name='export.pdf', mimetype='application/pdf')

    elif format == 'docx':
        doc = Document()
        doc.add_paragraph(text)
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return send_file(buffer, as_attachment=True, download_name='export.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    return jsonify({'error': 'Invalid format'}), 400
